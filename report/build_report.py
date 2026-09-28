"""Build the University of Mysore project report from the measured metrics."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from ml_pipeline import COLUMN_SPECS, EXCLUDED_FROM_MODEL, decode_frame, load_raw_frame  # noqa: E402
from report.diagrams import draw_all  # noqa: E402
from report.expand import extras, extra_tests, hardware_rows, literature_blocks  # noqa: E402

RESULTS = ROOT / "results" / "metrics.json"
OUT = ROOT / "Project_Report.docx"
FIG = ROOT / "report" / "figures"


def pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def set_run_font(run, size=12, bold=False, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_text(paragraph, text, size=12, bold=False, font="Times New Roman"):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, font=font)
    return run


def paragraph(doc, text, size=12, bold=False, center=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    add_text(p, text, size=size, bold=bold)
    return p


def chapter_heading(doc, label, title):
    heading(doc, label, 16, center=False)
    heading(doc, title, 18, center=True)


def write_blocks(doc, blocks):
    for text in blocks:
        if text.startswith("H2 "):
            heading(doc, text[3:], 16, center=False)
        elif text.startswith("H3 "):
            heading(doc, text[3:], 14, center=False)
        else:
            paragraph(doc, text)


def heading(doc, text, size, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, text, size=size, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Inches(0.35)
    add_text(p, "•  " + text, size=12)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_text(cell.paragraphs[0], header, size=11, bold=True)
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            cell.text = ""
            add_text(cell.paragraphs[0], str(value), size=11)
    doc.add_paragraph()
    return table


def add_picture(doc, path: Path, width=6.0):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def caption(doc, text):
    paragraph(doc, text, size=11, bold=True, center=True, space_after=12)


def page_break(doc):
    doc.add_page_break()


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Student Performance Prediction Using Machine Learning  ·  ", size=10)
    run = footer.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def build_declaration() -> None:
    doc = Document()
    configure(doc)
    heading(doc, "SELF DECLARATION OF ORIGINALITY AND ACADEMIC INTEGRITY", 16, center=True)
    paragraph(doc, "I, Siddhant Krushnasa Wadekar, Roll No. BCA22773, student of VI Semester BCA (Bachelor of Computer Applications), Directorate of Online Programme, University of Mysore, am submitting the project titled Student Performance Prediction Using Machine Learning for the degree examinations 2026–2027.")
    paragraph(doc, "I declare that the report and the program were prepared by me for this submission. The questionnaire table is the public UCI file named in the bibliography. The marks table is the sample file described in the report. The result tables are the output of my program. Papers and books I used are numbered in section 8.4 of the report. I have not copied another student's project.")
    paragraph(doc, "The department plagiarism check is still to be run by me before I sign. This page does not state a plagiarism percentage. I will sign only after that check, and only if the result is within the limit given in the project guidelines.")
    bullet(doc, "Signature of the student: _______________________________")
    bullet(doc, "Name: Siddhant Krushnasa Wadekar")
    bullet(doc, "Roll No.: BCA22773")
    bullet(doc, "Programme: VI Semester BCA, Directorate of Online Programme")
    bullet(doc, "Place: Yeola, Nashik, Maharashtra")
    bullet(doc, "Date: ____________________")
    path = ROOT / "Self_Declaration.docx"
    doc.save(path)
    print(path)


def add_frequency(doc) -> None:
    frame = decode_frame(load_raw_frame())
    heading(doc, "7.10 Counts of each answer", 16)
    paragraph(doc, "Table 7.4 lists every decoded answer and how many of the 145 students chose it. I used these counts while reading the file. A rare answer makes a one-hot column that is 1 on only a few rows. I did not delete rare answers, because they are real replies in the published file.")
    rows = []
    for column in frame.columns:
        if column in ("Student_ID", "Grade", "Performance_Category"):
            continue
        counts = frame[column].value_counts()
        used = "No" if column in EXCLUDED_FROM_MODEL or column == "Course_ID" else "Yes"
        for value, count in counts.items():
            rows.append([column.replace("_", " "), str(value), str(int(count)), f"{count / len(frame) * 100:.1f}%", used])
    add_table(doc, ["Column", "Answer", "Students", "Share", "In the model"], rows)
    caption(doc, "Table 7.4  Answer counts in the UCI file")


def add_student_table(doc) -> None:
    frame = decode_frame(load_raw_frame())
    heading(doc, "7.11 Rows used to train the model", 16)
    paragraph(doc, "Table 7.5 is the full set of 145 rows for the columns I looked at most often while checking the label. The grade is shown so that the category can be checked by hand. The grade is not an input to the model. Student id is only an identifier.")
    rows = []
    for _, row in frame.iterrows():
        rows.append(
            [
                str(row["Student_ID"]),
                str(row["Student_Age"]),
                str(row["Weekly_Study_Hours"]),
                str(row["Class_Attendance"]),
                str(row["Taking_Notes"]),
                str(row["Grade"]),
                str(row["Performance_Category"]),
            ]
        )
    add_table(
        doc,
        ["ID", "Age", "Study hours", "Attendance", "Notes", "Grade", "Category"],
        rows,
    )
    caption(doc, "Table 7.5  All 145 training rows, with the label shown for checking")


def add_marks_results(doc) -> None:
    from ml_pipeline import evaluate_marks

    report = evaluate_marks()
    heading(doc, "7.6.1 Sample marks model", 14)
    counts = report["class_counts"]
    baseline = report["majority_baseline_cv_accuracy"]["mean"]
    best = report["best_model"]
    best_row = report["models"][best]["cv"]
    paragraph(
        doc,
        "The sample has "
        f"{report['n_rows']} students: {counts['Excellent']} Excellent, "
        f"{counts['Average']} Average and {counts['Poor']} Poor. "
        "I built the exam percent from the three marks with some noise, then trained the same five algorithms. "
        "The published score is five-fold macro-F1. The exam percent is the label, so it is not an input. "
        f"Always answering {report['majority_class']} has accuracy {pct(baseline)}. "
        f"{best} has macro-F1 {pct(best_row['f1_macro']['mean'])} and accuracy {pct(best_row['accuracy']['mean'])}. "
        "On My result, this model supplies the predicted category. The rank uses the exam percent.",
    )
    rows = []
    for name, metrics in report["models"].items():
        cv = metrics["cv"]
        rows.append(
            [
                name + (" (selected)" if name == best else ""),
                pct(cv["accuracy"]["mean"]),
                pct(cv["f1_macro"]["mean"]),
                pct(cv["poor_recall"]["mean"]),
            ]
        )
    add_table(doc, ["Algorithm", "CV accuracy", "CV macro-F1", "Poor recall"], rows)
    caption(doc, "Table 7.6  Five-fold results for internal marks, assignment and attendance")


def add_arithmetic(doc, report, best) -> None:
    matrix = report["models"][best]["confusion_matrix"]
    labels = ["Excellent", "Average", "Poor"]
    total = sum(sum(row) for row in matrix)
    correct = sum(matrix[i][i] for i in range(3))
    heading(doc, "7.3.1 Manual check of the holdout arithmetic", 14)
    paragraph(
        doc,
        f"The holdout has {total} students. I add the diagonal of Table 7.2, which is "
        f"{matrix[0][0]} + {matrix[1][1]} + {matrix[2][2]} = {correct}. "
        f"Accuracy on this one split is {correct} divided by {total}, which is {correct / total * 100:.1f} percent. "
        "This is the holdout accuracy. It is not the five-fold accuracy in Table 7.1. "
        "I have left both numbers in the report because they answer different questions. "
        "The five-fold number is the one I used to choose the model. The holdout number is only this confusion matrix.",
    )
    for index, label in enumerate(labels):
        support = sum(matrix[index])
        predicted = sum(matrix[row][index] for row in range(3))
        hit = matrix[index][index]
        recall = (hit / support * 100) if support else 0
        precision = (hit / predicted * 100) if predicted else 0
        paragraph(
            doc,
            f"For the {label} row, the actual count is {support}. The model got {hit} of them right, so recall is {hit} / {support} = {recall:.1f} percent. "
            f"The {label} column sums to {predicted}, which is how many times the model said {label}. Precision is {hit} / {predicted} = {precision:.1f} percent. "
            "Precision and recall are not the same, so I do not quote only one of them. "
            "Macro recall is the average of the three recall values, and macro precision is the average of the three precision values. "
            "Macro-F1 uses both. A high score on Average alone cannot make macro-F1 high if Excellent or Poor is missed.",
        )
    paragraph(
        doc,
        "I also compared this holdout accuracy with the majority baseline. The baseline always says Average, so on a test set it is correct only on the Average row. "
        f"The Average row has {sum(matrix[1])} students out of {total}. That is the best a constant Average guess can do on this particular split, "
        "and it is a different number from the cross-validation baseline in Table 7.1, because the folds are not this same 29-row slice. "
        "I have not replaced the cross-validation baseline with this slice.",
    )


def add_website_screens(doc) -> None:
    heading(doc, "5.18 Screens from the website", 16)
    paragraph(doc, "The pictures below are taken from the website on this computer. The login names on the first screen are the demo accounts for the viva. They are not real college passwords.")
    shots = [
        ("screen_login.png", "Fig. 5.1  Login page"),
        ("screen_dashboard.png", "Fig. 5.2  Student dashboard, with the class counts and the selected model"),
        ("screen_predict.png", "Fig. 5.3  One questionnaire submission. The category is from the SVM, not from the exam."),
        ("screen_exam.png", "Fig. 5.4  Student exam, read from the question bank"),
        ("screen_my_result.png", "Fig. 5.5  Exam result with the actual band, the predicted category and the rank."),
        ("screen_questions.png", "Fig. 5.6  Question bank on the teacher login"),
        ("screen_internals.png", "Fig. 5.7  Internal marks saved by the teacher. This row is not used to train the model."),
        ("screen_evaluation.png", "Fig. 5.8  Evaluation page. The percentages match Table 7.1."),
    ]
    for filename, label in shots:
        path = FIG / filename
        if not path.exists():
            paragraph(doc, f"{label} was not found in the figures folder when this report was built.")
            continue
        add_picture(doc, path, 6.1)
        caption(doc, label)


def model_rows(report: dict) -> list[list[str]]:
    rows = []
    for name, metrics in report["models"].items():
        cv = metrics["cv"]
        mark = "Yes" if name == report["best_model"] else ""
        rows.append(
            [
                name,
                pct(cv["accuracy"]["mean"]),
                pct(cv["precision_macro"]["mean"]),
                pct(cv["recall_macro"]["mean"]),
                pct(cv["f1_macro"]["mean"]),
                pct(cv["f1_macro"]["std"]),
                pct(cv["poor_recall"]["mean"]),
                mark,
            ]
        )
    return rows


def build() -> None:
    report = json.loads(RESULTS.read_text())
    FIG.mkdir(parents=True, exist_ok=True)
    draw_all(FIG)
    build_declaration()
    best = report["best_model"]
    best_cv = report["models"][best]["cv"]
    counts = report["class_counts"]
    total = report["n_rows"]
    baseline = report["majority_baseline_cv_accuracy"]["mean"]

    doc = Document()
    configure(doc)

    for line, size, bold in [
        ("UNIVERSITY OF MYSORE", 16, True),
        ("Manasagangothri, Mysore – 570006", 14, False),
        ("", 12, False),
        ("STUDENT PERFORMANCE PREDICTION USING MACHINE LEARNING", 16, True),
        ("", 12, False),
        ("by", 14, False),
        ("", 12, False),
        ("Siddhant Krushnasa Wadekar", 14, True),
        ("", 12, False),
        ("VI Semester BCA", 14, True),
        ("Bachelor of Computer Applications", 14, False),
        ("", 12, False),
        ("Roll No. BCA22773", 14, True),
        ("", 12, False),
        ("Project Report submitted to the University of Mysore in partial fulfillment of the requirements of VI Semester BCA degree examinations 2026–2027", 14, False),
        ("", 12, False),
        ("University of Mysore", 14, True),
        ("Manasagangothri, Mysore – 570006", 14, False),
    ]:
        paragraph(doc, line, size=size, bold=bold, center=True, space_after=6)

    page_break(doc)
    heading(doc, "CERTIFICATE", 18, center=True)
    paragraph(doc, "This is to certify that the project report entitled “Student Performance Prediction Using Machine Learning” is prepared by Siddhant Krushnasa Wadekar, Reg. No. BCA22773, towards partial fulfillment of the requirements of the BCA degree of the University of Mysore.")

    heading(doc, "ACKNOWLEDGEMENT", 18, center=True)
    paragraph(doc, "I am thankful to the University of Mysore and the Directorate of Online Programme for giving me the chance to do this project. I also thank the authors of the UCI dataset, because the student records used in this project are taken from there.")

    heading(doc, "ABSTRACT", 18, center=True)
    paragraph(doc, f"This project is about predicting student performance with machine learning. The aim is to classify a student as Excellent, Average or Poor. I have used a dataset from UCI which has {total} students. There is no empty value in the file. Excellent means grade AA or BA ({counts['Excellent']} students). Average means BB, CB, CC or DC ({counts['Average']} students). Poor means DD or Fail ({counts['Poor']} students).")
    paragraph(doc, "The actual grade is not given as an input to the model, because that would make the prediction too easy. Last semester CGPA, expected CGPA, course id and sex are also not used as inputs. Course id was removed because the model was learning the course name instead of the student habits. Sex was removed because it is not useful for a teacher who wants to help a student in studies.")
    paragraph(doc, f"I trained five algorithms in Python using scikit-learn. They are Decision Tree, Logistic Regression, Random Forest, Naive Bayes and SVM. The model with the best average macro-F1 is selected. {best} came first. Its macro-F1 is {pct(best_cv['f1_macro']['mean'])} and its accuracy is {pct(best_cv['accuracy']['mean'])}. If we simply mark every student as Average, the accuracy is already {pct(baseline)}. So my model is not better than that simple guess. I have written this clearly in the results chapter instead of hiding it.")
    paragraph(doc, "When I added course id back just for checking, Decision Tree macro-F1 went up to "
               f"{pct(report['sensitivity']['plus_course_id']['Decision Tree']['f1_macro_mean'])}. "
               "That increase is not used in the final website, because it is not a fair way to predict. The website is made in Flask and it runs on the local computer for the project demo.")

    page_break(doc)
    heading(doc, "TABLE OF CONTENTS", 18, center=True)
    for item in [
        "Chapter I  Introduction",
        "Chapter II  Literature Survey",
        "Chapter III  System Requirements Specification",
        "Chapter IV  System Design",
        "Chapter V  Implementation",
        "Chapter VI  Testing",
        "Chapter VII  Results and Analysis",
        "Chapter VIII  Conclusion, Future Work and Bibliography",
        "Appendix A  Data dictionary",
        "Appendix B  How to run the program",
        "Appendix C  Source code",
    ]:
        paragraph(doc, item, center=False, space_after=2)

    heading(doc, "LIST OF TABLES", 16)
    for item in [
        "Table 2.1  How the cited work is used",
        "Table 3.1  Functional requirements",
        "Table 3.2  Software requirements",
        "Table 3.3  Hardware requirements",
        "Table 5.1  Columns excluded from the model",
        "Table 6.1  Tests executed on 26 September 2026",
        "Table 7.1  Five-fold results for the deployed feature set",
        "Table 7.2  Holdout confusion matrix for the selected model",
        "Table 7.3  Sensitivity to course ID and last-semester CGPA",
        "Table 7.4  Answer counts in the UCI file",
        "Table 7.5  Student rows used for training",
        "Table 7.6  Marks model: internal, assignment and attendance",
    ]:
        paragraph(doc, item, center=False, space_after=2)

    heading(doc, "LIST OF FIGURES", 16)
    for item in [
        "Fig. 4.1  Processing steps",
        "Fig. 4.2  Context data-flow diagram",
        "Fig. 4.3  Level-1 data-flow diagram",
        "Fig. 4.4  Use-case diagram",
        "Fig. 4.5  Class diagram",
        "Fig. 4.6  Sequence diagram",
        "Fig. 4.7  Activity diagram",
        "Fig. 4.8  E-R diagram",
        "Fig. 5.1  Login page",
        "Fig. 5.2  Student dashboard",
        "Fig. 5.3  Questionnaire prediction",
        "Fig. 5.4  Student exam",
        "Fig. 5.5  Exam result",
        "Fig. 5.6  Question bank",
        "Fig. 5.7  Internal marks",
        "Fig. 5.8  Model evaluation page",
        "Fig. 7.1  Students in each category",
        "Fig. 7.2  Macro-F1 by algorithm",
        "Fig. 7.3  Holdout confusion matrix",
    ]:
        paragraph(doc, item, center=False, space_after=2)

    heading(doc, "ABBREVIATIONS", 16)
    for item in [
        "BCA — Bachelor of Computer Applications",
        "CGPA — Cumulative grade point average",
        "CV — Cross-validation",
        "EDM — Educational data mining",
        "F1 — Harmonic mean of precision and recall",
        "ML — Machine learning",
        "SVM — Support vector machine",
        "UCI — University of California, Irvine Machine Learning Repository",
    ]:
        paragraph(doc, item, center=False, space_after=2)

    page_break(doc)
    chapter_heading(doc, "Chapter I", "INTRODUCTION")
    heading(doc, "1.1 Introduction to the project", 16)
    paragraph(doc, "In most colleges, a teacher gets to know that a student is weak only after the exam result is declared. By that time it is already late to help the student. Machine learning can be used to look at older student information and guess the performance category in advance. This idea comes under educational data mining [1], [2].")
    paragraph(doc, "For this project I did not collect data from my own college. I used a public dataset in which students have answered questions like age group, type of school, scholarship, job, study hours, attendance, notes and listening in class. From the final grade I made three groups, Excellent, Average and Poor, because these are the three categories in my project topic.")
    heading(doc, "1.2 Problem statement", 16)
    paragraph(doc, "The problem is to check whether study habits and background details can predict Excellent, Average or Poor without using the grade itself. If the accuracy is low, that result should also be shown. It is not correct to add a column like course id only to increase the percentage.")
    heading(doc, "1.3 Objectives", 16)
    bullet(doc, "To study the UCI student dataset and note down what each column means.")
    bullet(doc, "To divide grades into Excellent, Average and Poor.")
    bullet(doc, "To train Decision Tree, Logistic Regression, Random Forest, Naive Bayes and SVM.")
    bullet(doc, "To compare them using macro-F1 and also show accuracy against a simple baseline.")
    bullet(doc, "To make a small website with Admin, Teacher and Student login for the demo.")
    heading(doc, "1.4 Scope", 16)
    paragraph(doc, "This project uses the UCI file for the questionnaire model, and a separate sample marks file for internal marks, assignment and attendance. The exam percent is the actual result. It is not one of the three inputs. The website runs on my laptop for the viva. I have not connected it to a college database. The three login passwords are only for demonstration in the lab.")
    heading(doc, "1.5 Significance", 16)
    paragraph(doc, "The main learning from this project is that a high accuracy can be misleading. When course id was included, the score looked better, but the model was only remembering which course the student studied. After removing it, the accuracy became lower. Showing this lower score is more correct.")
    heading(doc, "1.6 Report outline", 16)
    paragraph(doc, "Chapter II explains the related work and the five algorithms. Chapter III gives the requirements. Chapter IV explains the design and the diagrams. Chapter V explains how I implemented the project, including the question bank and internal marks. Chapter VI shows the testing. Chapter VII shows the output and graphs. Chapter VIII gives the conclusion, the future work and the bibliography.")
    write_blocks(doc, extras("1"))

    page_break(doc)
    chapter_heading(doc, "Chapter II", "LITERATURE SURVEY")
    heading(doc, "2.1 System study", 16)
    paragraph(doc, "To predict a class, we need old records, a label, and some input columns. The label in my project is Excellent, Average or Poor. If I also keep the grade column as an input, the model will just copy the grade. That is not prediction. Han, Kamber and Pei also say that the class label should be kept separate from the input attributes [3].")
    heading(doc, "2.2 Educational data mining", 16)
    paragraph(doc, "Romero and Ventura have reviewed educational data mining. They say it includes prediction, clustering and finding relations in student data [1]. Baker and Yacef have also reviewed this area and mentioned that many people use public education datasets [2]. Both papers say that predicting pass or fail is a common task, but the inputs should be available before the result if the teacher wants to take action.")
    heading(doc, "2.3 A similar study", 16)
    paragraph(doc, "Cortez and Silva predicted secondary school marks in Mathematics and Portuguese using student background and study time [4]. They compared more than one algorithm on a small dataset. I have followed a similar method. I also compared more than one algorithm and I have not reported only accuracy.")
    heading(doc, "2.4 Dataset used by other authors", 16)
    paragraph(doc, "Yilmaz and Sekeroglu collected 145 student records from engineering and education students and tried to classify end-of-term performance [5]. The same file is available on the UCI website [6]. I downloaded that file and used it as it is for the questionnaire model. The UCI file has no internal-mark column, so the college module uses a separate sample marks file, described in section 7.6.1.")
    heading(doc, "2.5 Algorithms", 16)
    paragraph(doc, "Decision tree asks questions like if attendance is always, go left, otherwise go right, and at the end it gives a class. Quinlan explained how a tree can be built from examples [7]. On a small dataset a deep tree can memorise the training rows, so I have kept maximum depth as 5 and minimum samples in a leaf as 2.")
    paragraph(doc, "Logistic regression finds a relation between the input columns and the class. It is a basic model, so I kept it for comparison. After one-hot encoding there are many columns and only 145 rows, so I used the regularised version from scikit-learn.")
    paragraph(doc, "Random forest builds many decision trees and then takes the majority answer. Breiman showed that this reduces the mistake of a single tree [8]. I used 300 trees and maximum depth 8. Class weight is set to balanced so that Poor and Excellent are not ignored just because Average has more students.")
    paragraph(doc, "Naive Bayes calculates probability of each class from the columns. It assumes that the columns are independent. In this dataset that is not fully true, because one question is split into many yes/no columns. So a low score for Naive Bayes is possible.")
    paragraph(doc, "SVM tries to draw a boundary between the classes. Cortes and Vapnik explained the soft-margin SVM [9]. I used the RBF kernel. Before training, all columns are scaled, otherwise some columns can dominate the others.")
    heading(doc, "2.6 Comparison of literature", 16)
    paragraph(doc, "From the above papers I understood three things. Student marks can be predicted from student details. The best algorithm is not the same for every dataset. The result should mention which dataset and which split was used. Some papers use previous marks as input. My project does not use the grade as input.")
    heading(doc, "2.7 Gap", 16)
    paragraph(doc, "I am not proposing a new algorithm. I only checked, on this UCI file, whether study habit columns can still separate Excellent, Average and Poor after removing course id, sex and both CGPA columns. I also checked what happens if course id is added back.")
    write_blocks(doc, extras("2"))
    heading(doc, "2.10 Comparison table", 16)
    add_table(
        doc,
        ["Source", "What I used from it", "What I did not copy"],
        [
            ["Romero and Ventura [1]", "Prediction is one EDM task", "Their datasets"],
            ["Baker and Yacef [2]", "Report more than accuracy", "Their 2009 counts"],
            ["Han, Kamber and Pei [3]", "Keep the label out of the inputs", "Their textbook examples"],
            ["Cortez and Silva [4]", "Compare more than one model", "Their school marks"],
            ["Yilmaz and Sekeroglu [5], [6]", "The 145-row UCI file", "Their published accuracy"],
        ],
    )
    caption(doc, "Table 2.1  How the cited work is used in this project")
    write_blocks(doc, literature_blocks())

    page_break(doc)
    chapter_heading(doc, "Chapter III", "SYSTEM REQUIREMENTS SPECIFICATION")
    heading(doc, "3.1 Functional requirements", 16)
    paragraph(doc, "Table 3.1 shows what the program does. The UCI file has no question paper, so the exam, the question bank and the internal marks are stored in separate CSV files. Those files are not used to train the model.")
    add_table(
        doc,
        ["ID", "Requirement", "Role"],
        [
            ["FR1", "Log in as Admin, Teacher, or Student", "All"],
            ["FR2", "Reject an unknown password", "All"],
            ["FR3", "Show cohort counts and the selected model", "All"],
            ["FR4", "Predict a category from a completed profile", "All"],
            ["FR5", "Hide the student table from the Student login", "Student"],
            ["FR6", "Show the student table and the metric table", "Admin, Teacher"],
            ["FR7", "State the accuracy beside the majority baseline", "All"],
            ["FR8", "Add and remove questions", "Teacher"],
            ["FR9", "Submit the exam and see a percentage band", "Student"],
            ["FR10", "Enter internal marks, assignment and attendance", "Teacher"],
            ["FR11", "Add and remove a teacher name", "Admin"],
            ["FR12", "Keep the exam band separate from the SVM category", "All"],
        ],
    )
    caption(doc, "Table 3.1  Functional requirements")
    heading(doc, "3.2 Non-functional requirements", 16)
    bullet(doc, "random_state is kept as 42, so when I run the program again I get the same tables.")
    bullet(doc, "Training finishes in a few seconds because there are only 145 rows.")
    bullet(doc, "The website opens on 127.0.0.1, so it is only for this computer.")
    bullet(doc, "The predicted category is shown as a model result. It is not the official college mark.")
    heading(doc, "3.3 Hardware requirements", 16)
    paragraph(doc, "I ran this project on a normal laptop. Table 3.3 follows the university minimum. A graphics card is not required.")
    add_table(doc, ["Item", "Requirement"], hardware_rows())
    caption(doc, "Table 3.3  Hardware requirements")
    heading(doc, "3.4 Software requirements", 16)
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["Python 3", "Language"],
            ["pandas", "To read the CSV file"],
            ["scikit-learn", "To train and compare the models [10]"],
            ["matplotlib", "To draw the graphs"],
            ["Flask", "To make the login website"],
            ["python-docx", "To prepare this report"],
        ],
    )
    caption(doc, "Table 3.2  Software requirements")
    heading(doc, "3.5 Data requirements", 16)
    paragraph(doc, f"The file name is data/uci_student_performance.csv. It has {total} rows and 33 columns, including student id and grade. There is no missing value. The numbers in the file are codes, so I converted them into words using the UCI description. If some code is not in that list, the program stops instead of guessing the meaning.")
    heading(doc, "3.6 User requirements", 16)
    paragraph(doc, "During viva, the examiner can login as admin, teacher and student. Teacher and admin can see the student list and the accuracy table. Student login cannot open the full student list. There is no need to create a separate account for each student because this is a public dataset and not my college data.")
    heading(doc, "3.7 Feasibility", 16)
    paragraph(doc, "The project is technically possible because Python, Flask and scikit-learn are freely available and the data is small. It is not feasible to use this model in a real college for declaring results, because the accuracy is less than the simple baseline. It is suitable as a student project and demo.")
    write_blocks(doc, extras("3"))

    page_break(doc)
    chapter_heading(doc, "Chapter IV", "SYSTEM DESIGN")
    heading(doc, "4.1 Design overview", 16)
    paragraph(doc, "First the program reads the CSV and converts the codes into words. Then it removes the columns which should not be used. After that it trains five models and picks one. The scores written in Chapter VII are from cross validation, not from testing on the same rows used for training. For the website form, the selected model is trained once on all 145 rows when the program starts.")
    heading(doc, "4.2 System architecture", 16)
    add_picture(doc, FIG / "architecture.png", 6.2)
    caption(doc, "Fig. 4.1  Processing steps from the CSV to the model comparison")
    paragraph(doc, "The website does not train the model again on every click. Training happens once at the start. I have not used MySQL or any other database. The CSV file itself is the data store.")
    heading(doc, "4.3 Data flow diagram", 16)
    paragraph(doc, "Fig. 4.2 is the context diagram. The people and the files sit outside the circle. Fig. 4.3 is the level-1 diagram. Training reads the UCI store. The exam process reads and writes the question store. Internal marks write a different store. The grade column is used only to create the label. It is not passed as an input feature.")
    add_picture(doc, FIG / "dfd_context.png", 6.2)
    caption(doc, "Fig. 4.2  Context data-flow diagram")
    add_picture(doc, FIG / "dfd_level1.png", 6.2)
    caption(doc, "Fig. 4.3  Level-1 data-flow diagram")
    heading(doc, "4.4 Use case diagram", 16)
    add_picture(doc, FIG / "usecase.png", 6.2)
    caption(doc, "Fig. 4.4  Use-case diagram")
    paragraph(doc, "Admin and Teacher can see the student table and the model comparison. The teacher maintains questions and internal marks. The student takes the exam and fills the prediction form. If a student tries to open the student list page, the program sends him back. Hiding the menu link is not the only check. The page itself also checks the role.")
    heading(doc, "4.5 Class diagram", 16)
    add_picture(doc, FIG / "class_diagram.png", 6.2)
    caption(doc, "Fig. 4.5  Class diagram")
    paragraph(doc, "ml_pipeline.py reads the data, trains the models and saves the graphs. app.py has the login and the web pages. college_records.py reads and writes the question, mark and attempt files. Inside scikit-learn, one pipeline does one-hot encoding, then scaling, then the selected algorithm.")
    heading(doc, "4.6 Sequence diagram", 16)
    add_picture(doc, FIG / "sequence.png", 6.2)
    caption(doc, "Fig. 4.6  Sequence diagram for submitting the exam")
    paragraph(doc, "The student opens the exam page. Flask reads the question file and returns the form. After submit, Flask compares each chosen letter with the stored answer, writes the attempt, and shows the percentage category. This sequence does not call the SVM. The prediction form is a different page. On that page Flask checks the login, builds one row from the allowed columns, and the fitted model returns Excellent, Average or Poor.")
    heading(doc, "4.7 Activity diagram", 16)
    add_picture(doc, FIG / "activity.png", 6.2)
    caption(doc, "Fig. 4.7  Activity diagram after login")
    paragraph(doc, "The training activity, which is not the same as the login activity, is as follows. Read the CSV. Convert each code into text. Create the Excellent, Average and Poor label from the grade. Remove the columns which are not allowed. For every algorithm, split the data into 5 folds, train on 4 folds and test on 1 fold, and repeat this for all folds. Take the average macro-F1. Select the algorithm with the highest average. Train that algorithm on all rows for the website. Save the result file and the charts.")
    heading(doc, "4.8 E-R diagram", 16)
    add_picture(doc, FIG / "er.png", 6.2)
    caption(doc, "Fig. 4.8  E-R diagram")
    paragraph(doc, "The UCI student row is one entity. The exam attempt is another entity. They are not linked by a foreign key, because the survey students are not the demo logins. Internal marks are entered by name on the teacher page. The teacher list is a small table for the admin screen. Login passwords stay in the program for the demo.")
    heading(doc, "4.9 Modules", 16)
    bullet(doc, "Data module: read the UCI CSV and convert codes into words.")
    bullet(doc, "Model module: train five algorithms and compare them.")
    bullet(doc, "Exam module: question bank, student attempt and percentage band.")
    bullet(doc, "Internal-mark module: teacher entry of marks, assignment and attendance.")
    bullet(doc, "Web module: login, dashboard, prediction, student list and evaluation page.")
    write_blocks(doc, extras("4"))

    page_break(doc)
    chapter_heading(doc, "Chapter V", "IMPLEMENTATION")
    heading(doc, "5.1 Environment", 16)
    paragraph(doc, "I have written the project in Python. pandas is used to read the CSV. scikit-learn is used for the models. Flask is used for the website. matplotlib is used for the graphs. The steps to run the program are given in Appendix B.")
    heading(doc, "5.2 Data collection", 16)
    paragraph(doc, "I did not conduct a new survey. I downloaded the UCI file and saved it as data/uci_student_performance.csv. The UCI page says the data was collected in 2019 from engineering and education students [6]. The first questions are personal details, the next questions are about family, and the later questions are about study habits. Course id and grade are at the end.")
    heading(doc, "5.3 Preprocessing", 16)
    paragraph(doc, "The CSV stores answers as numbers. For example 1 may mean bus and 3 may mean bicycle. If I train the model on these numbers directly, it will think bicycle is greater than bus, which is wrong. So I first changed the numbers into words, and then used one-hot encoding. One-hot encoding makes a separate yes/no column for each answer. The file has no blank cell. Still I kept an imputer in the pipeline so that the program does not crash if a blank comes later.")
    heading(doc, "5.4 Feature selection", 16)
    paragraph(doc, "I did not use every column. The final model uses 27 columns such as age group, school type, scholarship, part-time work, sports, study hours, reading, attendance, notes, listening in class and similar questions. The columns which I removed are given in Table 5.1. I also ran the program again after adding course id and CGPA, only to see the difference. Those extra runs are not used in the website.")
    add_table(
        doc,
        ["Excluded column", "Reason"],
        [
            ["Student ID", "It is only an id, not a study habit"],
            ["Grade", "This is the answer we are predicting"],
            ["Performance category", "This is the same answer in three groups"],
            ["Last-semester CGPA", "This is already a result, not an input"],
            ["Expected graduation CGPA", "Student is already guessing his own result"],
            ["Course ID", "Model remembers the course. See Table 7.3"],
            ["Sex", "Not useful for helping the student in studies"],
        ],
    )
    caption(doc, "Table 5.1  Columns excluded from the model")
    heading(doc, "5.5 Model training and the selection rule", 16)
    paragraph(doc, "For every algorithm the steps are the same. First one-hot encoding, then standard scaling, then the algorithm. I used class_weight as balanced for Decision Tree, Random Forest, Logistic Regression and SVM, so the model does not only learn to say Average. Naive Bayes in scikit-learn does not have class_weight.")
    paragraph(doc, "I selected the model using average macro-F1 of 5 folds, with random_state 42. Macro-F1 gives equal importance to all three classes. I did not select the model using accuracy alone, because Average is almost half of the data and accuracy can look okay even if the model is weak. If two models have the same macro-F1, I would pick the one which finds more Poor students. I also kept one 80/20 test split for the confusion matrix. That test part has 29 students, so it is only to see the mistakes. It is not used to choose the winner.")
    heading(doc, "5.6 Prediction workflow", 16)
    paragraph(doc, "The prediction page shows dropdowns only for the columns used by the questionnaire model. After submit, it shows Excellent, Average or Poor. It also shows the percentage of each class when the model can give it. The rank of one student against the others is on My result, and it uses the exam percent, not this questionnaire score.")
    heading(doc, "5.7 Roles in the website", 16)
    paragraph(doc, "Admin and Teacher can see accuracy, precision, recall, macro-F1 and Poor recall for all five models. They can also see the 145 students with id, attendance, study hours, grade and category. The teacher can see predicted and actual on the internal-marks page. Student login can see the prediction form, the exam, and My result with the actual band, the predicted category and the rank. Student cannot open the evaluation page or the full list.")
    heading(doc, "5.8 Implementation issues", 16)
    bullet(doc, "After one-hot encoding, the number of columns becomes more than the number of students. So the model can memorise. I limited the tree depth to reduce this.")
    bullet(doc, "Each test fold has only about 29 students, so the score changes from fold to fold. That is why I have also written the standard deviation.")
    bullet(doc, "Family job, salary and parents education are still in the questionnaire model. A Poor label from these columns should not be announced as a final mark. The rank on My result uses the exam percent only.")
    bullet(doc, "Logistic regression gave some warning during training because many one-hot columns are related. I scaled the columns before training. Even after that, its accuracy is below the baseline.")
    heading(doc, "5.9 Security", 16)
    paragraph(doc, "The three passwords are written in the program so that login can be shown easily during viva. This is fine only on localhost. I have not used encryption or a real user database.")
    write_blocks(doc, extras("5"))
    add_website_screens(doc)

    page_break(doc)
    chapter_heading(doc, "Chapter VI", "TESTING")
    heading(doc, "6.1 Test environment", 16)
    paragraph(doc, "I tested the program on 26 September 2026. The command was python -m unittest tests.test_pipeline -v. It showed 15 tests and all were OK. Table 6.1 lists the checks. I have marked a row as Pass only if that check was actually run.")
    heading(doc, "6.2 Strategy", 16)
    paragraph(doc, "Some tests check the data, like number of rows, grade groups and removed columns. Some tests open the website pages using Flask test client. I checked wrong password, student trying to open the student list, teacher opening the evaluation page, and one prediction using the first student row.")
    heading(doc, "6.3 Results of the executed tests", 16)
    add_table(
        doc,
        ["ID", "Check", "Expected", "Result"],
        [
            ["T1", "Row count", "145 students", "Pass"],
            ["T2", "Labels", "Only Excellent, Average, Poor", "Pass"],
            ["T3", "Banding", "AA excellent, DC average, Fail poor", "Pass"],
            ["T4", "Exclusions", "Grade, both CGPA fields, course, sex absent", "Pass"],
            ["T5", "Models", "Five named algorithms and one selected model", "Pass"],
            ["T6", "Bad password", "Login page shows an error", "Pass"],
            ["T7", "Student opens /students", "Redirect and a refusal message", "Pass"],
            ["T8", "Teacher opens /evaluation", "Macro-F1 table is shown", "Pass"],
            ["T9", "Prediction of a known row", "Response contains one of the three labels", "Pass"],
            ["T10", "Dictionary", "CGPA fields are documented as excluded", "Pass"],
            *extra_tests(),
        ],
    )
    caption(doc, "Table 6.1  Tests executed on 26 September 2026")
    heading(doc, "6.4 Integration and system checks", 16)
    paragraph(doc, "The website and Chapter VII use the same result. The evaluation page shows the cross-validation average. The dashboard shows the selected model and the baseline accuracy. Prediction uses the model which was selected and then trained on all rows. This was checked through the teacher evaluation test and the prediction test.")
    heading(doc, "6.5 Defects", 16)
    paragraph(doc, "All 15 tests passed. The program is not crashing. The weak point is the questionnaire prediction in Chapter VII. The selected model is less accurate than always saying Average. I have not changed the test to expect a high accuracy.")
    heading(doc, "6.6 What was not tested", 16)
    paragraph(doc, "I checked the pages with the Flask test client. A full click-through in the browser at http://127.0.0.1:8080 can still be done before the viva. I did not do load testing because there are only 145 records.")
    write_blocks(doc, extras("6"))

    page_break(doc)
    chapter_heading(doc, "Chapter VII", "RESULTS AND ANALYSIS")
    heading(doc, "7.1 Cohort", 16)
    paragraph(doc, f"Excellent has {counts['Excellent']} students ({pct(counts['Excellent']/total)}). Average has {counts['Average']} students ({pct(counts['Average']/total)}). Poor has {counts['Poor']} students ({pct(counts['Poor']/total)}). Average is the biggest group. If the model does nothing and always answers Average, accuracy comes to about {pct(baseline)}.")
    add_picture(doc, ROOT / "results" / "class_distribution.png", 5.6)
    caption(doc, "Fig. 7.1  Students in each performance category (n = 145)")
    heading(doc, "7.2 Model comparison", 16)
    paragraph(doc, f"Table 7.1 shows the main result. {best} has the highest average macro-F1, which is {pct(best_cv['f1_macro']['mean'])}. The variation between folds is {pct(best_cv['f1_macro']['std'])}. Its accuracy is {pct(best_cv['accuracy']['mean'])}. This is less than the baseline of {pct(baseline)}. Recall for the Poor class is {pct(best_cv['poor_recall']['mean'])}. So even the best model is not finding Poor students properly, and it is also less correct than simply marking everyone as Average.")
    add_table(
        doc,
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro-F1", "F1 std", "Poor recall", "Selected"],
        model_rows(report),
    )
    caption(doc, "Table 7.1  Five-fold means for the deployed feature set")
    add_picture(doc, ROOT / "results" / "model_comparison.png", 5.8)
    caption(doc, "Fig. 7.2  Mean macro-F1 by algorithm, with the fold standard deviation")
    heading(doc, "7.3 Holdout errors", 16)
    paragraph(doc, f"Table 7.2 is the confusion matrix of {best} on the 20 percent test data, which has 29 students. This one split will not match the 5-fold average exactly. Rows are the actual category and columns are the predicted category. The order is Excellent, Average, Poor.")
    matrix = report["models"][best]["confusion_matrix"]
    add_table(
        doc,
        ["Actual \\ Predicted", "Excellent", "Average", "Poor"],
        [
            ["Excellent", *matrix[0]],
            ["Average", *matrix[1]],
            ["Poor", *matrix[2]],
        ],
    )
    caption(doc, f"Table 7.2  Holdout confusion matrix for {best} (29 students)")
    add_picture(doc, ROOT / "results" / "confusion_matrix.png", 5.2)
    caption(doc, f"Fig. 7.3  Holdout confusion matrix for {best}")
    paragraph(doc, "From the matrix we can see that many students are predicted in the wrong group. If a college uses this output, some Poor students will be missed and some Average students will be marked as Poor. This is why the macro-F1 is low.")
    add_arithmetic(doc, report, best)
    heading(doc, "7.4 What raises the score for the wrong reason", 16)
    paragraph(doc, "Table 7.3 shows what happens if I add back one removed column. When course id is added, Decision Tree macro-F1 goes above 50 percent and accuracy becomes better than the baseline. This happens because the model learns which course is easy or hard, not because it understood study habits. So I did not keep course id in the final model. When last semester CGPA is added, the score increases only a little and accuracy is still below the baseline. I still removed CGPA because it is already a result.")
    sens_rows = []
    for setting, title in [
        ("deployed_features", "Deployed columns"),
        ("plus_course_id", "Deployed plus course ID"),
        ("plus_last_semester_cgpa", "Deployed plus last-semester CGPA"),
    ]:
        block = report["sensitivity"][setting]
        winner = max(block, key=lambda name: block[name]["f1_macro_mean"])
        sens_rows.append(
            [
                title,
                winner,
                pct(block[winner]["f1_macro_mean"]),
                pct(block[winner]["accuracy_mean"]),
            ]
        )
    add_table(
        doc,
        ["Feature set", "Best by macro-F1", "Macro-F1", "Accuracy"],
        sens_rows,
    )
    caption(doc, "Table 7.3  Sensitivity to course ID and last-semester CGPA")
    heading(doc, "7.5 Tree features, as a description only", 16)
    paragraph(doc, "The final model is SVM, and SVM does not give a simple list of important columns. So I also checked the decision tree on the same columns, only for explanation. The tree gave more weight to mother education, scholarship, age group 18-21, additional work, father education and parental status. These are background details. It does not mean that changing one habit will move a student from Poor to Average. The tree accuracy is also not better than the baseline in a useful way, as shown in Table 7.1.")
    heading(doc, "7.6 Discussion", 16)
    paragraph(doc, "After removing the shortcut columns, these questions are not enough to predict the three groups on 145 students. One reason can be that the dataset is small. Another reason can be that final grade depends more on the exam than on these survey answers. Table 7.3 shows that course id is the column which changes the result the most.")
    paragraph(doc, "The synopsis asked for internal marks, assignment, attendance, an exam, a predicted category and a rank. Those columns are not in the UCI file, so I did not type them into the 145 UCI rows. I prepared a separate sample of 90 students in data/college_marks.csv. The inputs are internal marks, assignment and attendance. The label is the exam category. Table 7.6 is that comparison. The UCI tables above are a different model.")
    add_marks_results(doc)
    heading(doc, "7.7 Limits", 16)
    bullet(doc, "Only one dataset is used and it has 145 rows.")
    bullet(doc, "I used one random seed. If the seed is changed, the percentage can move a little.")
    bullet(doc, "Some columns are family and income details, which the student cannot change easily.")
    bullet(doc, "The website trains the selected model on all 145 rows for the form. The tables in this chapter are from cross validation, not from that final training.")
    write_blocks(doc, extras("7"))
    add_frequency(doc)
    add_student_table(doc)

    page_break(doc)
    chapter_heading(doc, "Chapter VIII", "CONCLUSION AND FUTURE WORK")
    heading(doc, "8.1 Conclusion", 16)
    paragraph(doc, f"In this project I prepared a Python program and a small website to predict student performance as Excellent, Average or Poor. Five algorithms were compared on the UCI dataset. {best} got the best macro-F1, so it is selected. Its accuracy is {pct(best_cv['accuracy']['mean'])}, which is less than {pct(baseline)} accuracy of always predicting Average. So this project should not be used to decide a student's real result. It can be used to understand how the models behave on this dataset.")
    heading(doc, "8.2 Contributions", 16)
    bullet(doc, "Converted the UCI number codes into readable answers.")
    bullet(doc, "Removed grade, both CGPA columns, course id and sex from the model.")
    bullet(doc, "Compared five algorithms and also showed a simple baseline.")
    bullet(doc, "Made a demo website with Admin, Teacher and Student login, an exam, a question bank, internal marks, a predicted category and a rank.")
    heading(doc, "8.3 Future work", 16)
    paragraph(doc, "In future, the sample marks file can be replaced with records from my college, with permission. The same three inputs can stay: internal marks, assignment and attendance. The questionnaire model still needs a larger file before it can be used for a real result. Tuning the parameters will be useful only when that file is bigger.")
    heading(doc, "8.4 Bibliography", 16)
    references = [
        "[1] Romero, C. and Ventura, S. (2010). Educational data mining: A review of the state of the art. IEEE Transactions on Systems, Man, and Cybernetics, Part C, 40(6), 601–618.",
        "[2] Baker, R. S. J. D. and Yacef, K. (2009). The state of educational data mining in 2009: A review and future visions. Journal of Educational Data Mining, 1(1), 3–17.",
        "[3] Han, J., Kamber, M. and Pei, J. (2012). Data Mining: Concepts and Techniques. 3rd ed. Morgan Kaufmann.",
        "[4] Cortez, P. and Silva, A. (2008). Using data mining to predict secondary school student performance. In Proceedings of the 5th Future Business Technology Conference. EUROSIS.",
        "[5] Yılmaz, N. and Şekeroğlu, B. (2020). Student performance classification using artificial intelligence techniques. In Advances in Intelligent Systems and Computing, vol. 1095. Springer.",
        "[6] Yılmaz, N. and Şekeroğlu, B. (2019). Higher Education Students Performance Evaluation. UCI Machine Learning Repository. https://doi.org/10.24432/C51G82",
        "[7] Quinlan, J. R. (1986). Induction of decision trees. Machine Learning, 1(1), 81–106.",
        "[8] Breiman, L. (2001). Random forests. Machine Learning, 45(1), 5–32.",
        "[9] Cortes, C. and Vapnik, V. (1995). Support-vector networks. Machine Learning, 20(3), 273–297.",
        "[10] Pedregosa, F. et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    ]
    for ref in references:
        paragraph(doc, ref, center=False)
    write_blocks(doc, extras("8"))

    page_break(doc)
    chapter_heading(doc, "Appendix A", "DATA DICTIONARY")
    paragraph(doc, "The table below shows the meaning of each code in the UCI file. In the program I stored the words, not the numbers. Grade codes are 0 Fail, 1 DD, 2 DC, 3 CC, 4 CB, 5 BB, 6 BA and 7 AA.")
    dict_rows = []
    for _source, name, mapping in COLUMN_SPECS:
        coded = "; ".join(f"{code} = {label}" for code, label in mapping.items())
        used = "No" if name in EXCLUDED_FROM_MODEL else "Yes"
        dict_rows.append([name.replace("_", " "), used, coded])
    dict_rows.append(["Course ID", "No", "Integer course identifier, shown as Course N"])
    dict_rows.append(["Grade", "No", "0 Fail; 1 DD; 2 DC; 3 CC; 4 CB; 5 BB; 6 BA; 7 AA"])
    add_table(doc, ["Column", "In the model", "Codes"], dict_rows)
    caption(doc, "Table A.1  Decoded columns and whether the model uses them")

    page_break(doc)
    chapter_heading(doc, "Appendix B", "HOW TO RUN THE PROGRAM")
    paragraph(doc, "Install the packages given in requirements.txt. Then open the project folder and run python -m unittest tests.test_pipeline -v. After that run python app.py. Open the browser and go to http://127.0.0.1:8080. Do not use port 5000, because on Mac that port is used by AirPlay and the browser shows 403. Login details are admin / admin123, teacher / teacher123 and student / student123. These passwords are only for the project demo.")
    paragraph(doc, "If the report has to be prepared again, first run the training file so that results/metrics.json is created, and then run python report/build_report.py.")

    page_break(doc)
    chapter_heading(doc, "Appendix C", "SOURCE CODE")
    paragraph(doc, "The source code of the project is given below. It is added in the report because the university guidelines say that the program should be included with the theory.")
    for relative in ["ml_pipeline.py", "app.py", "college_records.py", "tests/test_pipeline.py"]:
        heading(doc, relative, 14)
        source = (ROOT / relative).read_text()
        for line in source.splitlines():
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            add_text(p, line if line else " ", size=11, font="Courier New")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
